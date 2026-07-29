import os
import calendar
import yfinance as yf
import pandas as pd
from datetime import datetime, timedelta
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ticker = "NEX.PA"
excel_file = "Nexans_VWAP.xlsx"

# --- Détermination du mois cible ---
# Variable d'environnement MOIS au format "AAAA-MM" (ex: "2026-07")
# Si absente : mois précédent par défaut (utile pour l'exécution automatique en début de mois)
mois_param = os.environ.get("MOIS", "").strip()

if mois_param:
    annee, mois = map(int, mois_param.split("-"))
else:
    today = datetime.today()
    premier_jour_mois_actuel = today.replace(day=1)
    dernier_jour_mois_precedent = premier_jour_mois_actuel - timedelta(days=1)
    annee, mois = dernier_jour_mois_precedent.year, dernier_jour_mois_precedent.month

start_date = f"{annee}-{mois:02d}-01"
dernier_jour = calendar.monthrange(annee, mois)[1]
end_date_exclusive = (datetime(annee, mois, dernier_jour) + timedelta(days=1)).strftime("%Y-%m-%d")

print(f"Calcul du VWAP pour {annee}-{mois:02d} (du {start_date} au {dernier_jour:02d}/{mois:02d}/{annee})")

# --- Téléchargement des données du mois ---
data = yf.download(ticker, start=start_date, end=end_date_exclusive, interval="1d", auto_adjust=False, actions=False)

# Aplatir les colonnes si yfinance renvoie un MultiIndex (Prix, Ticker)
if isinstance(data.columns, pd.MultiIndex):
    data.columns = data.columns.droplevel(1)

if "Close" not in data.columns or data["Close"].dropna().empty:
    print("Aucune donnée disponible pour ce mois (marché fermé tout le mois, ticker invalide, ou mois futur).")
else:
    data.index = pd.to_datetime(data.index).tz_localize(None)

    # VWAP cumulatif jour par jour, à l'intérieur du mois uniquement
    data["VWAP"] = (data["Close"] * data["Volume"]).cumsum() / data["Volume"].cumsum()

    # VWAP du mois entier (une seule valeur) = somme(Close*Volume) / somme(Volume) sur tout le mois
    vwap_mensuel = (data["Close"] * data["Volume"]).sum() / data["Volume"].sum()
    vwap_mensuel = float(vwap_mensuel)

    sheet_mois = f"{annee}-{mois:02d}"

    # Charger le classeur existant s'il existe, sinon mode écriture simple
    mode = "a" if os.path.exists(excel_file) else "w"
    kwargs = {"if_sheet_exists": "replace"} if mode == "a" else {}

    with pd.ExcelWriter(excel_file, engine="openpyxl", mode=mode, **kwargs) as writer:
        # Onglet du mois avec tableau Excel filtrable
        data.to_excel(writer, sheet_name=sheet_mois, index=True, index_label="Date")
        ws = writer.sheets[sheet_mois]
        n_rows = data.shape[0] + 1
        n_cols = data.shape[1] + 1
        ref = f"A1:{get_column_letter(n_cols)}{n_rows}"
        tbl = Table(displayName=f"Table_{annee}{mois:02d}", ref=ref)
        style = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True, showColumnStripes=False)
        tbl.tableStyleInfo = style
        ws.add_table(tbl)

        # Mise à jour de l'onglet récapitulatif "VWAP_Mensuel" (une ligne par mois)
        if mode == "a" and "VWAP_Mensuel" in pd.ExcelFile(excel_file).sheet_names:
            recap = pd.read_excel(excel_file, sheet_name="VWAP_Mensuel")
        else:
            recap = pd.DataFrame(columns=["Mois", "VWAP"])

        recap = recap[recap["Mois"] != sheet_mois]  # évite les doublons si on relance le même mois
        recap = pd.concat([recap, pd.DataFrame([{"Mois": sheet_mois, "VWAP": vwap_mensuel}])], ignore_index=True)
        recap = recap.sort_values("Mois").reset_index(drop=True)
        recap.to_excel(writer, sheet_name="VWAP_Mensuel", index=False)

    print(f"Fichier Excel mis à jour : {excel_file}")
    print(f"VWAP de {sheet_mois} : {vwap_mensuel:.2f} €")
